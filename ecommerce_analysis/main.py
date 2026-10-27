from pathlib import Path
import math
import sys

import pandas as pd


# 获取项目目录，并优先查找指定名称的数据文件。
BASE_DIR = Path(__file__).resolve().parent
DATA_FILE_CANDIDATES = [
    BASE_DIR / "input" / "amazon_sales_data.csv",
    BASE_DIR / "amazon_sales_data.csv",
    # 兼容当前工作区中提供的 Amazon.csv 文件。
    BASE_DIR / "input" / "Amazon.csv",
    BASE_DIR / "input:" / "Amazon.csv",
]
OUTPUT_DIR = BASE_DIR / "output"
REPORT_FILE = OUTPUT_DIR / "data_quality_report.xlsx"
SALES_REPORT_FILE = OUTPUT_DIR / "sales_analysis_report.xlsx"


def main():
    """读取销售数据，输出基础信息并生成数据质量检查报告。"""
    data_file = next(
        (path for path in DATA_FILE_CANDIDATES if path.is_file()),
        None,
    )
    if data_file is None:
        expected_paths = "\n".join(str(path) for path in DATA_FILE_CANDIDATES)
        raise FileNotFoundError(
            f"找不到销售数据文件，请将 CSV 放在以下任一位置：\n{expected_paths}"
        )

    # 使用 pandas 读取 CSV 文件。
    data = pd.read_csv(data_file)
    print(f"数据文件：{data_file}")

    # 查看前 5 行数据。
    print("\n前 5 行数据：")
    print(data.head())

    # 输出列名、非空值数量、数据类型和内存占用等基本信息。
    print("\n数据基本信息：")
    data.info(buf=sys.stdout)

    # 统计每列缺失值的数量。
    print("\n各列缺失值统计：")
    print(data.isna().sum())

    # 输出数值列的描述性统计；include='all' 同时涵盖非数值列。
    print("\n描述性统计：")
    print(data.describe(include="all"))

    # 全部检查仅针对内存中的数据执行，不会修改或覆盖原始 CSV。
    report_sheets, quality_summary, prompts = build_quality_report(data)
    print("\n数据质量检查汇总：")
    print(quality_summary.to_string(index=False))
    if prompts:
        print("\n检查提示：")
        for prompt in prompts:
            print(f"- {prompt}")

    # 报告目录不存在时自动创建，Excel 文件只写入 output 目录。
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if REPORT_FILE.exists():
        print(f"\n已保留现有数据质量报告（本次不覆盖）：{REPORT_FILE}")
    else:
        write_formatted_excel(report_sheets, REPORT_FILE)
        print(f"\n数据质量报告已保存：{REPORT_FILE}")

    # 销售业务分析依据数据中的实际字段生成独立报告，不覆盖原始 CSV 或质量报告。
    sales_sheets, field_report, business_warnings = build_sales_analysis(data)
    print("\n字段识别结果：")
    print(field_report.to_string(index=False))
    print("\n业务分析提示：")
    for warning in business_warnings:
        print(f"- {warning}")

    # 已有报告不重建，避免覆盖其工作表；仅在文件尚不存在时创建基础分析。
    if not SALES_REPORT_FILE.exists():
        write_formatted_excel(sales_sheets, SALES_REPORT_FILE)
        print(f"\n销售业务分析报告已创建：{SALES_REPORT_FILE}")
    else:
        print(f"\n已保留现有销售分析工作表：{SALES_REPORT_FILE}")

    return_report, return_warnings = build_return_analysis(data)
    if return_warnings:
        print("\n退货订单分析提示：")
        for warning in return_warnings:
            print(f"- {warning}")
    update_excel_sheet(SALES_REPORT_FILE, "退货订单分析", return_report)
    print(f"\n退货订单分析已写入：{SALES_REPORT_FILE}")

    grouped_return_report, grouped_return_warnings = build_grouped_return_rates(data)
    if grouped_return_warnings:
        print("\n分组退货率分析提示：")
        for warning in grouped_return_warnings:
            print(f"- {warning}")
    update_grouped_return_rate_report(
        SALES_REPORT_FILE,
        grouped_return_report,
    )
    print(f"\n分组退货率已写入：{SALES_REPORT_FILE}")


def build_quality_report(data: pd.DataFrame) -> tuple[dict[str, pd.DataFrame], pd.DataFrame, list[str]]:
    """根据实际列类型生成质量检查工作表、汇总表和跳过提示。"""
    prompts = []

    # 缺失值统计：保留每一列，即使该列没有缺失值。
    missing_report = (
        data.isna()
        .sum()
        .rename("缺失值数量")
        .rename_axis("列名")
        .reset_index()
    )

    # duplicated 默认保留每组重复数据的第一条，只统计其余重复行，不删除任何记录。
    duplicate_count = int(data.duplicated(keep="first").sum())
    duplicate_report = pd.DataFrame(
        [
            {"检查项": "重复行数量（不含每组第一条）", "数量": duplicate_count},
            {"检查项": "整行数据总数", "数量": len(data)},
        ]
    )

    # 数据类型和后续数值、文本检查均从 CSV 实际推断出的列类型动态获取。
    dtype_report = pd.DataFrame(
        {
            "列名": data.columns,
            "数据类型": data.dtypes.astype(str).to_numpy(),
            "非空值数量": data.notna().sum().to_numpy(),
        }
    )
    numeric_columns = data.select_dtypes(include="number").columns.tolist()
    text_columns = data.select_dtypes(include=["object", "string"]).columns.tolist()

    # 负数逐数值列统计；负值仅作为待核查信息，不自动判定为错误。
    if numeric_columns:
        negative_report = pd.DataFrame(
            [
                {
                    "列名": column,
                    "负数数量": int(data[column].lt(0).sum()),
                    "说明": "仅报告负数，不据此判断数据错误。",
                }
                for column in numeric_columns
            ]
        )
    else:
        negative_report = pd.DataFrame(
            columns=["列名", "负数数量", "说明"]
        )
        prompts.append("未发现数值列，已跳过负数检查。")

    # 使用 IQR（四分位距）围栏识别数值列中的疑似极端值，只报告、不删除。
    outlier_rows = []
    for column in numeric_columns:
        values = data[column].dropna()
        if values.empty:
            outlier_rows.append(
                {
                    "列名": column,
                    "疑似异常值数量": 0,
                    "下界": None,
                    "上界": None,
                    "示例值（最多 10 个）": "",
                    "检查说明": "该列没有非空数值，无法计算 IQR。",
                }
            )
            continue

        first_quartile = values.quantile(0.25)
        third_quartile = values.quantile(0.75)
        interquartile_range = third_quartile - first_quartile
        lower_bound = first_quartile - 1.5 * interquartile_range
        upper_bound = third_quartile + 1.5 * interquartile_range
        outliers = values.loc[(values < lower_bound) | (values > upper_bound)]
        examples = outliers.head(10).tolist()
        outlier_rows.append(
            {
                "列名": column,
                "疑似异常值数量": int(len(outliers)),
                "下界": lower_bound,
                "上界": upper_bound,
                "示例值（最多 10 个）": ", ".join(map(str, examples)),
                "检查说明": "按 1.5 倍 IQR 围栏筛查，仅供人工核查。",
            }
        )

    if not numeric_columns:
        outlier_report = pd.DataFrame(
            columns=[
                "列名", "疑似异常值数量", "下界", "上界",
                "示例值（最多 10 个）", "检查说明",
            ]
        )
        prompts.append("未发现数值列，已跳过疑似异常值检查。")
    else:
        outlier_report = pd.DataFrame(outlier_rows)

    # 文本检查涵盖空字符串、首尾空格，以及仅大小写/首尾空格不同的值。
    text_quality_rows = []
    empty_string_total = 0
    surrounding_space_total = 0
    inconsistent_group_total = 0
    inconsistent_row_total = 0
    for column in text_columns:
        raw_values = data[column].dropna().astype(str)
        stripped_values = raw_values.str.strip()
        empty_count = int(stripped_values.eq("").sum())
        whitespace_count = int(raw_values.ne(stripped_values).sum())
        empty_string_total += empty_count
        surrounding_space_total += whitespace_count

        text_quality_rows.extend(
            [
                {
                    "列名": column,
                    "检查项": "空字符串（不含缺失值）",
                    "涉及数量": empty_count,
                    "示例": "",
                },
                {
                    "列名": column,
                    "检查项": "首尾存在空格",
                    "涉及数量": whitespace_count,
                    "示例": ", ".join(
                        raw_values.loc[raw_values.ne(stripped_values)]
                        .drop_duplicates()
                        .head(5)
                        .tolist()
                    ),
                },
            ]
        )

        normalized_values = stripped_values.str.casefold()
        variants = pd.DataFrame(
            {"原始值": raw_values, "标准化值": normalized_values}
        )
        variant_groups = variants.groupby("标准化值")["原始值"].agg(
            lambda values: values.drop_duplicates().tolist()
        )
        inconsistent_groups = variant_groups.loc[variant_groups.map(len) > 1]
        affected_rows = int(
            normalized_values.isin(inconsistent_groups.index).sum()
        )
        inconsistent_group_total += len(inconsistent_groups)
        inconsistent_row_total += affected_rows
        variant_examples = [
            " / ".join(values[:4])
            for values in inconsistent_groups.head(5).tolist()
        ]
        text_quality_rows.append(
            {
                "列名": column,
                "检查项": "仅大小写或首尾空格不同的格式变体",
                "涉及数量": (
                    f"{len(inconsistent_groups)} 组；{affected_rows} 行"
                ),
                "示例": "; ".join(variant_examples),
            }
        )

    if not text_columns:
        text_quality_rows.append(
            {
                "列名": "",
                "检查项": "跳过提示",
                "涉及数量": 0,
                "示例": "未发现文本列，已跳过文本格式检查。",
            }
        )
        prompts.append("未发现文本列，已跳过空字符串和文本格式检查。")

    text_quality_report = pd.DataFrame(text_quality_rows)
    if not text_columns:
        text_quality_report = text_quality_report.reindex(
            columns=["列名", "检查项", "涉及数量", "示例"]
        )

    # 汇总指标说明全表检查结果；疑似异常值按列计数，跨列可能涉及同一行。
    missing_cells = int(data.isna().sum().sum())
    negative_total = int(
        sum(int(data[column].lt(0).sum()) for column in numeric_columns)
    )
    outlier_total = int(
        sum(row["疑似异常值数量"] for row in outlier_rows)
    )
    quality_summary = pd.DataFrame(
        [
            {"检查指标": "数据行数", "结果": len(data), "说明": ""},
            {"检查指标": "数据列数", "结果": len(data.columns), "说明": ""},
            {
                "检查指标": "存在缺失值的列数",
                "结果": int(data.isna().any().sum()),
                "说明": "",
            },
            {"检查指标": "缺失值总数", "结果": missing_cells, "说明": ""},
            {
                "检查指标": "重复行数量（不含每组第一条）",
                "结果": duplicate_count,
                "说明": "仅统计，不删除。",
            },
            {
                "检查指标": "负数总数",
                "结果": negative_total,
                "说明": "按数值列累计；仅报告，不判定为错误。",
            },
            {
                "检查指标": "疑似异常值总数",
                "结果": outlier_total,
                "说明": "按列累计，可能重复涉及同一行；仅报告，不删除。",
            },
            {
                "检查指标": "空字符串总数",
                "结果": empty_string_total,
                "说明": "仅统计文本列，不含缺失值。",
            },
            {
                "检查指标": "首尾空格文本值总数",
                "结果": surrounding_space_total,
                "说明": "仅统计文本列。",
            },
            {
                "检查指标": "格式变体组数",
                "结果": inconsistent_group_total,
                "说明": "只检查忽略大小写和首尾空格后相同的值。",
            },
            {
                "检查指标": "格式变体涉及行数",
                "结果": inconsistent_row_total,
                "说明": "可能与其他文本质量问题重复。",
            },
        ]
    )

    if not len(data.columns):
        prompts.append("数据没有任何字段，无法执行逐列数据质量检查。")

    report_sheets = {
        "数据质量汇总": quality_summary,
        "缺失值统计": missing_report,
        "重复数据检查": duplicate_report,
        "负数检查": negative_report,
        "数据类型": dtype_report,
        "异常值检查": outlier_report,
        "文本格式检查": text_quality_report,
        "检查提示": pd.DataFrame({"提示": prompts or ["所有适用检查均已执行。"]}),
    }
    return report_sheets, quality_summary, prompts


def normalize_field_name(name: str) -> str:
    """标准化字段名以匹配明确的常见字段别名。"""
    return "".join(character for character in name.casefold() if character.isalnum())


def identify_business_fields(data: pd.DataFrame) -> tuple[dict[str, str], pd.DataFrame]:
    """按明确别名、数据类型和样例识别字段；未匹配字段不推断业务含义。"""
    role_aliases = {
        "日期": {"date", "orderdate", "transactiondate", "salesdate"},
        "订单编号": {"orderid", "ordernumber", "transactionid"},
        "产品/SKU": {"sku", "productid", "productsku", "itemid"},
        "产品名称": {"productname", "producttitle", "itemname"},
        "销量": {"quantity", "unitsold", "unitssold", "soldquantity"},
        "报告金额": {"totalamount", "salesamount", "revenue", "totalrevenue"},
        "商品单价": {"unitprice", "price", "sellingprice"},
        "国家": {"country"},
        "州/省": {"state", "province"},
        "城市": {"city"},
        "地区/区域": {"region", "salesregion", "territory"},
        "品类": {"category", "productcategory"},
        "品牌": {"brand"},
        "广告费用": {"adcost", "advertisingcost", "advertising spend", "adspend"},
        "成本": {"cost", "totalcost", "productcost", "unitcost"},
        "利润": {"profit", "netprofit", "grossprofit"},
        "退款金额": {"refundamount", "refund", "returnamount"},
        "折扣": {"discount"},
        "税费": {"tax"},
        "运费": {"shippingcost", "shippingfee"},
        "订单状态": {"orderstatus", "status"},
    }
    role_candidates: dict[str, list[str]] = {}
    field_roles: dict[str, list[str]] = {column: [] for column in data.columns}
    field_notes: dict[str, str] = {}

    for column in data.columns:
        normalized = normalize_field_name(str(column))
        roles = [
            role
            for role, aliases in role_aliases.items()
            if normalized in {normalize_field_name(alias) for alias in aliases}
        ]
        for role in roles:
            # 日期字段必须可解析；数量和金额须为数值型，避免仅凭列名误作分析。
            if role == "日期":
                parsed = pd.to_datetime(data[column], errors="coerce")
                valid_ratio = float(parsed.notna().mean()) if len(data) else 0.0
                if valid_ratio < 0.95:
                    field_notes[column] = (
                        f"字段名类似日期，但仅 {valid_ratio:.1%} 可解析；不用于趋势分析。"
                    )
                    continue
            if role in {"销量", "报告金额", "商品单价", "广告费用", "成本",
                        "利润", "退款金额", "折扣", "税费", "运费"}:
                if not pd.api.types.is_numeric_dtype(data[column]):
                    field_notes[column] = (
                        f"字段名类似“{role}”，但数据类型为 {data[column].dtype}；"
                        "不进行数值汇总。"
                    )
                    continue
            role_candidates.setdefault(role, []).append(column)
            field_roles[column].append(role)
            if role == "日期":
                field_notes[column] = f"日期可解析比例 {valid_ratio:.1%}。"

    # 同一业务角色出现多个候选列时不擅自选择，避免重复字段导致汇总口径不确定。
    recognized: dict[str, str] = {}
    for role, columns in role_candidates.items():
        if len(columns) == 1:
            recognized[role] = columns[0]
        else:
            for column in columns:
                field_roles[column] = [
                    existing for existing in field_roles[column]
                    if existing != role
                ]
                field_roles[column].append(f"{role}候选（多列歧义）")
                field_notes[column] = (
                    f"多个字段均可能表示“{role}”：{', '.join(columns)}；"
                    "为避免擅自选择，该角色不用于汇总。"
                )

    rows = []
    for column in data.columns:
        examples = data[column].dropna().head(3).tolist()
        roles = field_roles[column]
        note = field_notes.get(column, "")
        if not roles:
            roles_text = "未明确识别"
            if not note:
                note = "仅列出原始字段名、类型和样例；不擅自推断业务含义。"
        else:
            roles_text = "、".join(roles)
            if not note:
                note = "依据字段名的明确别名识别；业务口径仍需数据提供方确认。"
        rows.append(
            {
                "字段名称": column,
                "数据类型": str(data[column].dtype),
                "示例值（最多 3 个）": ", ".join(map(str, examples)),
                "识别角色": roles_text,
                "说明": note,
            }
        )
    return recognized, pd.DataFrame(rows)


def _sum_column(data: pd.DataFrame, column: str | None) -> float | int | None:
    """返回可用数值列的原始汇总值。"""
    if column is None:
        return None
    return data[column].sum()


def build_sales_analysis(
    data: pd.DataFrame,
) -> tuple[dict[str, pd.DataFrame], pd.DataFrame, list[str]]:
    """基于实际识别到的字段生成销售分析结果、字段字典和业务提示。"""
    recognized, field_report = identify_business_fields(data)
    warnings: list[str] = list(dict.fromkeys(
        str(row["说明"])
        for _, row in field_report.iterrows()
        if "多个字段均可能表示" in str(row["说明"])
    ))
    report_sheets: dict[str, pd.DataFrame] = {"字段识别": field_report}

    date_column = recognized.get("日期")
    order_column = recognized.get("订单编号")
    product_column = recognized.get("产品/SKU") or recognized.get("产品名称")
    quantity_column = recognized.get("销量")
    amount_column = recognized.get("报告金额")
    price_column = recognized.get("商品单价")

    # 订单数只在出现明确订单编号字段时计算，并另列缺失和重复编号情况。
    order_count = None
    if order_column is not None:
        order_values = data[order_column].dropna()
        order_count = int(order_values.nunique())
        missing_order_ids = int(data[order_column].isna().sum())
        repeated_order_rows = int(data[order_column].duplicated(keep="first").sum())
        if missing_order_ids:
            warnings.append(
                f"订单编号字段“{order_column}”有 {missing_order_ids} 条缺失记录；"
                "订单数仅按非空编号去重计算。"
            )
        if repeated_order_rows:
            warnings.append(
                f"订单编号字段“{order_column}”在多行重复（{repeated_order_rows} 行）；"
                "总金额相加可能混合订单级与商品行级口径，需确认数据粒度。"
            )
        elif len(order_values) == order_count:
            warnings.append(
                f"字段“{order_column}”每条记录均唯一，按非空编号去重得到订单数；"
                "该数据快照未显示多商品订单行，但数据粒度仍需业务确认。"
            )
    else:
        warnings.append("未识别到可靠的订单编号字段，无法计算订单数。")

    # 整体概览只汇总真实存在且已识别的数量、金额和产品标识字段。
    overview_rows = [
        {"指标": "记录数", "结果": int(len(data)), "单位/计算口径": "行"},
        {
            "指标": "产品数",
            "结果": (
                int(data[product_column].nunique(dropna=True))
                if product_column is not None else None
            ),
            "单位/计算口径": (
                f"{product_column} 去重数"
                if product_column is not None else "无明确产品标识或产品名称字段，无法计算"
            ),
        },
        {
            "指标": "订单数",
            "结果": order_count,
            "单位/计算口径": (
                f"{order_column} 非空去重数"
                if order_column is not None else "无明确订单编号字段，无法计算"
            ),
        },
        {
            "指标": "总销量",
            "结果": _sum_column(data, quantity_column),
            "单位/计算口径": (
                f"{quantity_column} 求和；单位按原字段，未额外推断"
                if quantity_column is not None else "无明确数值型销量字段，无法计算"
            ),
        },
        {
            "指标": "报告金额合计",
            "结果": _sum_column(data, amount_column),
            "单位/计算口径": (
                f"{amount_column} 求和；币种及收入/实收口径未确认"
                if amount_column is not None else "无明确数值型金额字段，无法计算"
            ),
        },
    ]
    report_sheets["整体概览"] = pd.DataFrame(overview_rows)
    if quantity_column is None:
        warnings.append("未识别到明确的数值型销量字段，无法汇总总销量。")
    if amount_column is None:
        warnings.append("未识别到明确的数值型金额字段，无法计算金额合计。")
    else:
        warnings.append(
            f"“{amount_column}”仅按字段值汇总；币种、是否包含税费/运费、"
            "是否为实收收入均未由数据字典确认。"
        )

    # 产品排名以已识别的产品标识分组，不把单价当作销售额。
    if product_column is not None:
        group_columns = [product_column]
        name_column = next(
            (
                column for column in data.columns
                if column != product_column
                and normalize_field_name(str(column)) == "productname"
            ),
            None,
        )
        if name_column is not None:
            inconsistent_names = int(
                (data.groupby(product_column, dropna=False)[name_column].nunique(dropna=True) > 1).sum()
            )
            if inconsistent_names:
                warnings.append(
                    f"有 {inconsistent_names} 个产品标识对应多个产品名称；"
                    "排名仍以“{product_column}”为准。"
                )
            group_columns.append(name_column)
        aggregations: dict[str, tuple[str, str]] = {
            "记录数": (product_column, "size"),
        }
        if quantity_column:
            aggregations["销量合计"] = (quantity_column, "sum")
        if amount_column:
            aggregations["报告金额合计"] = (amount_column, "sum")
        product_report = (
            data.groupby(group_columns, dropna=False)
            .agg(**aggregations)
            .reset_index()
        )
        ranking_column = "报告金额合计" if amount_column else (
            "销量合计" if quantity_column else "记录数"
        )
        product_report = product_report.sort_values(
            ranking_column, ascending=False, kind="stable"
        ).head(20)
        report_sheets["产品排名"] = product_report
    else:
        warnings.append("未识别到明确的产品标识或产品名称字段，跳过产品表现排名。")

    # 日期通过解析比例门槛后按月汇总；无可靠日期时不生成趋势表。
    if date_column is not None:
        parsed_dates = pd.to_datetime(data[date_column], errors="coerce")
        valid_rows = data.loc[parsed_dates.notna()].copy()
        valid_rows["分析月份"] = parsed_dates.loc[parsed_dates.notna()].dt.to_period("M").astype(str)
        aggregations = {"记录数": (date_column, "size")}
        if quantity_column:
            aggregations["销量合计"] = (quantity_column, "sum")
        if amount_column:
            aggregations["报告金额合计"] = (amount_column, "sum")
        if order_column:
            aggregations["订单数"] = (order_column, "nunique")
        trend = (
            valid_rows.groupby("分析月份", dropna=False)
            .agg(**aggregations)
            .reset_index()
            .sort_values("分析月份")
        )
        report_sheets["销售趋势"] = trend
        if len(valid_rows) != len(data):
            warnings.append(
                f"日期字段“{date_column}”有 {len(data) - len(valid_rows)} 条无法解析，"
                "趋势仅包含可解析记录。"
            )
    else:
        warnings.append("未找到可可靠解析的日期字段，跳过按月销售趋势。")

    # 按明确地理字段分别输出地区表现，避免把行政层级混合在同一维度。
    region_columns = [
        (role, column)
        for role in ("country", "region", "state", "province", "city")
        for column in data.columns
        if normalize_field_name(str(column)) == role
    ]
    if region_columns:
        region_frames = []
        for level, column in region_columns:
            aggregations = {"记录数": (column, "size")}
            if order_column:
                aggregations["订单数"] = (order_column, "nunique")
            if quantity_column:
                aggregations["销量合计"] = (quantity_column, "sum")
            if amount_column:
                aggregations["报告金额合计"] = (amount_column, "sum")
            grouped = (
                data.groupby(column, dropna=False)
                .agg(**aggregations)
                .reset_index()
                .rename(columns={column: "地区"})
            )
            grouped.insert(0, "地区字段", column)
            grouped.insert(1, "字段层级", level)
            region_frames.append(grouped)
        report_sheets["地区分析"] = pd.concat(region_frames, ignore_index=True)
    else:
        warnings.append("未识别到明确的地区字段，跳过地区分析。")

    # 品类与品牌分开汇总；只有明确存在的分类字段才生成相应工作表。
    for role, normalized_name, sheet_name in (
        ("品类", "category", "品类分析"),
        ("品牌", "brand", "品牌分析"),
    ):
        category_column = next(
            (
                column for column in data.columns
                if normalize_field_name(str(column)) == normalized_name
            ),
            None,
        )
        if category_column is None:
            warnings.append(f"未识别到明确的{role}字段，跳过{role}分析。")
            continue
        aggregations = {"记录数": (category_column, "size")}
        if order_column:
            aggregations["订单数"] = (order_column, "nunique")
        if quantity_column:
            aggregations["销量合计"] = (quantity_column, "sum")
        if amount_column:
            aggregations["报告金额合计"] = (amount_column, "sum")
        report_sheets[sheet_name] = (
            data.groupby(category_column, dropna=False)
            .agg(**aggregations)
            .reset_index()
            .sort_values(
                "报告金额合计" if amount_column else "记录数",
                ascending=False,
                kind="stable",
            )
        )

    # 价格分析仅描述明确单价字段的原始分布，不将单价误作金额或收入。
    if price_column is not None:
        prices = data[price_column].dropna()
        price_stats = prices.describe(percentiles=[0.25, 0.5, 0.75])
        report_sheets["价格分析"] = pd.DataFrame(
            [
                {"统计项": "非空记录数", "结果": int(price_stats["count"]), "单位/口径": f"{price_column}"},
                {"统计项": "平均值", "结果": price_stats["mean"], "单位/口径": f"{price_column} 原始单位未确认"},
                {"统计项": "中位数", "结果": price_stats["50%"], "单位/口径": f"{price_column} 原始单位未确认"},
                {"统计项": "标准差", "结果": price_stats["std"], "单位/口径": f"{price_column} 原始单位未确认"},
                {"统计项": "最小值", "结果": price_stats["min"], "单位/口径": f"{price_column} 原始单位未确认"},
                {"统计项": "第 25 百分位", "结果": price_stats["25%"], "单位/口径": f"{price_column} 原始单位未确认"},
                {"统计项": "第 75 百分位", "结果": price_stats["75%"], "单位/口径": f"{price_column} 原始单位未确认"},
                {"统计项": "最大值", "结果": price_stats["max"], "单位/口径": f"{price_column} 原始单位未确认"},
            ]
        )
    else:
        warnings.append("未识别到明确的数值型单价字段，跳过价格分布分析。")

    # 订单状态可以描述记录分布；退货/取消金额不能仅靠状态字段推算。
    status_column = next(
        (
            column for column in data.columns
            if normalize_field_name(str(column)) == "orderstatus"
        ),
        None,
    )
    if status_column is not None:
        aggregations = {"记录数": (status_column, "size")}
        if order_column:
            aggregations["订单数"] = (order_column, "nunique")
        if quantity_column:
            aggregations["销量合计"] = (quantity_column, "sum")
        if amount_column:
            aggregations["报告金额合计"] = (amount_column, "sum")
        report_sheets["订单状态分析"] = (
            data.groupby(status_column, dropna=False)
            .agg(**aggregations)
            .reset_index()
            .sort_values("记录数", ascending=False)
        )
        status_examples = ", ".join(map(str, data[status_column].dropna().unique()[:10]))
        warnings.append(
            f"状态字段“{status_column}”的值包括：{status_examples}。"
            "总量与分组金额均未按状态过滤；退货/取消记录是否计入销售额需业务确认。"
        )

    # 若存在明确的广告、成本、利润或退款金额字段，报告其原始汇总，不臆造利润公式。
    financial_roles = ("广告费用", "成本", "利润", "退款金额")
    financial_rows = []
    for role in financial_roles:
        column = recognized.get(role)
        if column is None:
            metric_name = role
            warnings.append(f"未识别到明确的{role}数值字段，无法分析{metric_name}。")
            continue
        financial_rows.append(
            {
                "指标类别": role,
                "字段名称": column,
                "非空记录数": int(data[column].notna().sum()),
                "字段值合计": data[column].sum(),
                "负数记录数": int(data[column].lt(0).sum()),
                "说明": "仅汇总该字段原始值；币种、统计粒度及业务口径需确认。",
            }
        )
    if financial_rows:
        report_sheets["成本广告退款分析"] = pd.DataFrame(financial_rows)
    else:
        warnings.append(
            "当前数据没有明确的广告费、成本、利润或退款金额数值字段；"
            "无法计算广告投入、成本、利润或退款金额。"
        )

    # 指标说明记录公式、单位边界和必要的解释限制。
    metric_rows = [
        {"指标/分析": "记录数", "计算公式": "数据行数", "单位": "行", "适用条件与限制": "不是订单数；记录粒度未完全确认。"},
        {"指标/分析": "产品数", "计算公式": "明确产品/SKU字段的非空去重数", "单位": "个", "适用条件与限制": "按字段标识统计，产品标识是否稳定需数据字典确认。"},
        {"指标/分析": "订单数", "计算公式": "明确订单编号字段的非空去重数", "单位": "个订单编号", "适用条件与限制": "只在识别到可靠订单号时计算；不以记录数代替。"},
        {"指标/分析": "总销量", "计算公式": "明确数值型销量字段求和", "单位": "原字段单位", "适用条件与限制": "字段须明确表示件数/销量；不进行单位换算。"},
        {"指标/分析": "报告金额合计", "计算公式": "明确数值型金额字段求和", "单位": "原始币种未确认", "适用条件与限制": "不是经验证的净销售额/实收金额；核对状态、税费、运费及粒度。"},
        {"指标/分析": "产品/地区/品类表现", "计算公式": "按明确维度分组并分别求记录数、订单数、销量、报告金额", "单位": "字段原单位", "适用条件与限制": "各指标仅在相应字段存在时输出；重复订单行可能造成金额口径风险。"},
        {"指标/分析": "月度趋势", "计算公式": "按可解析日期的年月分组汇总", "单位": "月", "适用条件与限制": "无可靠日期时跳过；不可解析日期记录不纳入趋势。"},
        {"指标/分析": "价格分布", "计算公式": "单价字段的均值、中位数、分位数、极值", "单位": "单价原始单位", "适用条件与限制": "只描述单价分布，不代表销售额。"},
        {"指标/分析": "广告/成本/利润/退款", "计算公式": "存在相应明确数值字段时汇总原始字段值", "单位": "原始币种未确认", "适用条件与限制": "无字段则不计算；不以销售额减成本等方式推造利润。"},
    ]
    report_sheets["指标说明"] = pd.DataFrame(metric_rows)
    report_sheets["业务提示"] = pd.DataFrame(
        [{"提示": warning} for warning in warnings]
        or [{"提示": "未发现需要提示的字段缺失或口径限制。"}]
    )
    return report_sheets, field_report, warnings


def find_single_column(
    data: pd.DataFrame,
    aliases: set[str],
) -> tuple[str | None, list[str]]:
    """只在字段名与明确别名唯一匹配时返回字段，避免多列歧义。"""
    normalized_aliases = {normalize_field_name(alias) for alias in aliases}
    matches = [
        column
        for column in data.columns
        if normalize_field_name(str(column)) in normalized_aliases
    ]
    return (matches[0] if len(matches) == 1 else None), matches


def _numeric_summary_rows(
    data: pd.DataFrame,
    column: str,
    returned_rows: pd.DataFrame,
    delivered_rows: pd.DataFrame,
    section_name: str,
    warnings: list[str],
    flag_nonpositive: bool = False,
) -> list[dict]:
    """生成分组数值描述统计，并报告缺失、非数值及可疑取值。"""
    results = {}
    for group_name, group_data in (
        ("退货组 Returned", returned_rows),
        ("正常交付组 Delivered", delivered_rows),
    ):
        raw = group_data[column]
        numeric = pd.to_numeric(raw, errors="coerce")
        nonfinite_mask = numeric.notna() & ~numeric.map(
            lambda value: bool(pd.notna(value) and abs(value) != float("inf"))
        )
        finite_values = numeric.loc[~nonfinite_mask].dropna()
        invalid_count = int((raw.notna() & numeric.isna()).sum())
        missing_count = int(raw.isna().sum())
        nonfinite_count = int(nonfinite_mask.sum())
        stats = finite_values.describe(percentiles=[0.25, 0.5, 0.75])
        if stats["count"]:
            stat_values = {
                key: (
                    float(stats[key])
                    if pd.notna(stats[key]) and math.isfinite(float(stats[key]))
                    else None
                )
                for key in ("mean", "50%", "std", "min", "25%", "75%", "max")
            }
        else:
            stat_values = {
                key: None
                for key in ("mean", "50%", "std", "min", "25%", "75%", "max")
            }
        results[group_name] = {
            "有效样本数": int(finite_values.count()),
            "缺失值数": missing_count,
            "非数值数": invalid_count,
            "非有限值数": nonfinite_count,
            "均值": stat_values["mean"],
            "中位数": stat_values["50%"],
            "标准差": stat_values["std"],
            "最小值": stat_values["min"],
            "25%分位数": stat_values["25%"],
            "75%分位数": stat_values["75%"],
            "最大值": stat_values["max"],
            "非正数/可疑值数": (
                int((finite_values <= 0).sum()) if flag_nonpositive else None
            ),
        }

    rows: list[dict] = []
    stat_names = (
        "有效样本数", "缺失值数", "非数值数", "非有限值数",
        "均值", "中位数", "标准差", "最小值", "25%分位数",
        "75%分位数", "最大值", "非正数/可疑值数",
    )
    for stat_name in stat_names:
        if stat_name == "非正数/可疑值数" and not flag_nonpositive:
            continue
        returned_value = results["退货组 Returned"][stat_name]
        delivered_value = results["正常交付组 Delivered"][stat_name]
        if stat_name in {"有效样本数", "缺失值数", "非数值数", "非有限值数",
                         "非正数/可疑值数"}:
            difference = None
        elif returned_value is not None and delivered_value is not None:
            difference = returned_value - delivered_value
        else:
            difference = None
        rows.append(
            {
                "分析分区": section_name,
                "指标或类别": stat_name,
                "退货组 Returned": returned_value,
                "正常交付组 Delivered": delivered_value,
                "差异/说明": difference,
            }
        )

    for group_name, group_data in (
        ("退货组 Returned", returned_rows),
        ("正常交付组 Delivered", delivered_rows),
    ):
        raw = group_data[column]
        numeric = pd.to_numeric(raw, errors="coerce")
        invalid_count = int((raw.notna() & numeric.isna()).sum())
        missing_count = int(raw.isna().sum())
        nonfinite_count = int(
            (numeric.notna() & numeric.map(
                lambda value: bool(pd.notna(value) and abs(value) == float("inf"))
            )).sum()
        )
        if missing_count or invalid_count or nonfinite_count:
            warnings.append(
                f"{column}（{group_name}）存在 {missing_count} 个缺失、"
                f"{invalid_count} 个非数值、{nonfinite_count} 个非有限值；"
                "描述统计仅使用可解析的有限数值。"
            )
        if flag_nonpositive:
            nonpositive_values = numeric.loc[numeric.notna() & (numeric <= 0)]
            nonpositive_count = int(nonpositive_values.count())
            if nonpositive_count:
                warnings.append(
                    f"{column}（{group_name}）有 {nonpositive_count} 个非正数，"
                    f"示例：{nonpositive_values.head(5).tolist()}；"
                    "仅标记供核查，未剔除。"
                )
    return rows


def build_return_analysis(
    data: pd.DataFrame,
) -> tuple[pd.DataFrame, list[str]]:
    """按明确的 Returned / Delivered 状态生成退货订单描述性对比。"""
    warnings: list[str] = []
    output_rows: list[dict] = []
    output_columns = [
        "分析分区",
        "指标或类别",
        "退货组 Returned",
        "正常交付组 Delivered",
        "差异/说明",
    ]

    def add_section(title: str) -> None:
        output_rows.append(
            {
                "分析分区": title,
                "指标或类别": None,
                "退货组 Returned": None,
                "正常交付组 Delivered": None,
                "差异/说明": None,
            }
        )

    def add_row(
        section: str,
        metric: str,
        returned_value=None,
        delivered_value=None,
        note=None,
    ) -> None:
        output_rows.append(
            {
                "分析分区": section,
                "指标或类别": metric,
                "退货组 Returned": returned_value,
                "正常交付组 Delivered": delivered_value,
                "差异/说明": note,
            }
        )

    status_column, status_candidates = find_single_column(
        data, {"orderstatus", "status"}
    )
    order_column, order_candidates = find_single_column(
        data, {"orderid", "ordernumber", "transactionid"}
    )
    price_column, price_candidates = find_single_column(
        data, {"unitprice", "sellingprice", "price"}
    )
    discount_column, discount_candidates = find_single_column(
        data,
        {"discount", "discountrate", "discountpercentage", "discountpercent"},
    )
    category_column, category_candidates = find_single_column(
        data, {"category", "productcategory"}
    )

    add_section("字段识别与状态筛选")
    fields = [
        ("订单状态字段", status_column, status_candidates),
        ("订单编号字段", order_column, order_candidates),
        ("单价字段", price_column, price_candidates),
        ("折扣字段", discount_column, discount_candidates),
        ("品类字段", category_column, category_candidates),
    ]
    for label, field, candidates in fields:
        if field is not None:
            add_row("字段识别与状态筛选", label, field, None, "按唯一明确字段名匹配。")
        elif candidates:
            note = f"候选字段不唯一：{', '.join(map(str, candidates))}；跳过该字段分析。"
            add_row("字段识别与状态筛选", label, None, None, note)
            warnings.append(note)
        else:
            add_row("字段识别与状态筛选", label, None, None, "字段不存在，相关分析跳过。")

    if status_column is None:
        warnings.append("未能唯一识别订单状态字段，无法筛选退货组或正常交付组。")
        add_section("样本概览")
        add_row("样本概览", "状态筛选", None, None, "缺少可靠订单状态字段，无法比较。")
        add_section("简要观察")
        add_row("简要观察", "观察", None, None, "没有可可靠识别的状态字段，无法进行组间分析。")
        return pd.DataFrame(output_rows, columns=output_columns), warnings

    normalized_status = data[status_column].map(
        lambda value: str(value).strip().casefold() if pd.notna(value) else ""
    )
    returned_mask = normalized_status.eq("returned")
    # 仅明确表示已交付的状态纳入对照；Shipped、Pending、Cancelled 等不合并。
    delivered_statuses = {"delivered", "successfully delivered"}
    delivered_mask = normalized_status.isin(delivered_statuses)
    returned_rows = data.loc[returned_mask]
    delivered_rows = data.loc[delivered_mask]
    excluded_statuses = sorted(
        {
            str(value).strip()
            for value in data.loc[
                ~(returned_mask | delivered_mask), status_column
            ].dropna().unique()
        }
    )

    add_row(
        "字段识别与状态筛选",
        "退货组规则",
        "状态经去除首尾空格、忽略大小写后等于 Returned",
        None,
        f"识别到 {len(returned_rows)} 条记录。",
    )
    add_row(
        "字段识别与状态筛选",
        "正常交付组规则",
        None,
        "状态经去除首尾空格、忽略大小写后等于 Delivered 或 Successfully Delivered",
        f"识别到 {len(delivered_rows)} 条记录；其他状态不纳入对照。",
    )
    add_row(
        "字段识别与状态筛选",
        "未纳入比较的其他状态",
        None,
        None,
        ", ".join(excluded_statuses) if excluded_statuses else "无",
    )

    add_section("样本概览")
    compared_records = len(returned_rows) + len(delivered_rows)
    returned_record_share = (
        len(returned_rows) / compared_records if compared_records else None
    )
    delivered_record_share = (
        len(delivered_rows) / compared_records if compared_records else None
    )
    add_row(
        "样本概览",
        "记录数",
        len(returned_rows),
        len(delivered_rows),
        f"对比组记录总数：{compared_records}。",
    )
    add_row(
        "样本概览",
        "记录占比（仅 Returned + Delivered）",
        returned_record_share,
        delivered_record_share,
        "分母为退货组与正常交付组记录数之和；其他状态排除。",
    )

    returned_order_count = None
    delivered_order_count = None
    if order_column is not None:
        returned_order_count = int(returned_rows[order_column].nunique(dropna=True))
        delivered_order_count = int(delivered_rows[order_column].nunique(dropna=True))
        returned_order_ids = set(returned_rows[order_column].dropna().tolist())
        delivered_order_ids = set(delivered_rows[order_column].dropna().tolist())
        total_distinct_orders = len(returned_order_ids | delivered_order_ids)
        returned_order_share = (
            returned_order_count / total_distinct_orders
            if total_distinct_orders else None
        )
        delivered_order_share = (
            delivered_order_count / total_distinct_orders
            if total_distinct_orders else None
        )
        add_row(
            "样本概览",
            "非空订单编号数量（去重）",
            returned_order_count,
            delivered_order_count,
            f"识别字段：{order_column}；不以记录数替代订单数。",
        )
        add_row(
            "样本概览",
            "订单编号占比（两组订单编号并集）",
            returned_order_share,
            delivered_order_share,
            "若同一编号同时出现在两种状态，组间占比会重叠。",
        )
        add_row(
            "样本概览",
            "缺失订单编号记录数",
            int(returned_rows[order_column].isna().sum()),
            int(delivered_rows[order_column].isna().sum()),
            "缺失编号无法计入去重订单数。",
        )
        add_row(
            "样本概览",
            "重复订单编号记录数（组内）",
            int(returned_rows[order_column].duplicated().sum()),
            int(delivered_rows[order_column].duplicated().sum()),
            "用于区分记录数与订单数；同一编号重复时记录可能是订单商品行。",
        )
        if returned_order_ids.intersection(delivered_order_ids):
            warnings.append(
                f"有 {len(returned_order_ids.intersection(delivered_order_ids))} 个订单编号"
                "同时出现在退货与交付组，订单组别存在重叠。"
            )
    else:
        add_row(
            "样本概览",
            "非空订单编号数量（去重）",
            None,
            None,
            "没有唯一明确订单编号字段，无法计算订单数或订单占比。",
        )
        warnings.append("未识别到唯一明确订单编号字段，订单数无法计算。")

    if data[status_column].isna().any():
        warnings.append(
            f"状态字段“{status_column}”有 {int(data[status_column].isna().sum())} 条缺失；"
            "这些记录不属于任一比较组。"
        )
    if excluded_statuses:
        warnings.append(
            "以下状态未纳入退货/正常交付对比：" + ", ".join(excluded_statuses) + "。"
        )
    if not len(returned_rows):
        warnings.append("未找到状态为 Returned 的记录，退货组分析为空。")
    if not len(delivered_rows):
        warnings.append(
            "未找到明确表示已交付的 Delivered 记录；不将其他状态当作正常交付，"
            "无法进行组间比较。"
        )

    add_section("单价分布")
    if price_column is not None and len(returned_rows) and len(delivered_rows):
        price_start = len(output_rows)
        output_rows.extend(
            _numeric_summary_rows(
                data,
                price_column,
                returned_rows,
                delivered_rows,
                "单价分布",
                warnings,
                flag_nonpositive=True,
            )
        )
        for row in output_rows[price_start:]:
            row["差异/说明"] = (
                f"退货组减正常交付组；字段“{price_column}”原始单位未确认。"
                if row["指标或类别"] in {"均值", "中位数", "标准差", "最小值",
                                          "25%分位数", "75%分位数", "最大值"}
                else (
                    "非正单价仅标记人工核查，未自动剔除。"
                    if row["指标或类别"] == "非正数/可疑值数"
                    else row["差异/说明"]
                )
            )
    else:
        reason = (
            "未能唯一识别明确的单价字段。"
            if price_column is None
            else "缺少任一比较组样本，无法进行组间单价统计。"
        )
        add_row("单价分布", "分析状态", None, None, reason)
        warnings.append(reason)

    add_section("折扣分布")
    if discount_column is not None and len(returned_rows) and len(delivered_rows):
        discount_start = len(output_rows)
        output_rows.extend(
            _numeric_summary_rows(
                data,
                discount_column,
                returned_rows,
                delivered_rows,
                "折扣分布",
                warnings,
            )
        )
        discount_values = pd.to_numeric(data[discount_column], errors="coerce").dropna()
        discount_values = discount_values.loc[
            discount_values.map(lambda value: abs(value) != float("inf"))
        ]
        if discount_values.empty:
            raw_value_note = (
                f"字段“{discount_column}”没有可用有限数值；"
                "折扣单位未确认，未换算为百分数或金额。"
            )
        else:
            unique_values = sorted(discount_values.unique().tolist())
            raw_value_note = (
                f"字段“{discount_column}”保留原始值；当前观察范围 "
                f"{discount_values.min():g} 至 {discount_values.max():g}，"
                f"不同取值 {len(unique_values)} 个。"
                "这些值可能采用比例口径，但未获得字段定义，未换算为百分数或折扣金额。"
            )
        for row in output_rows[discount_start:]:
            if row["指标或类别"] in {"均值", "中位数", "标准差", "最小值",
                                      "25%分位数", "75%分位数", "最大值"}:
                row["差异/说明"] = (
                    "退货组减正常交付组；" + raw_value_note
                )
    else:
        reason = (
            "未能唯一识别明确的折扣数值字段。"
            if discount_column is None
            else "缺少任一比较组样本，无法进行组间折扣统计。"
        )
        add_row("折扣分布", "分析状态", None, None, reason)
        warnings.append(reason)

    add_section("品类分布（按记录统计）")
    category_rows: list[dict] = []
    if category_column is not None and len(returned_rows) and len(delivered_rows):
        observed_categories = pd.concat(
            [returned_rows[category_column], delivered_rows[category_column]],
            ignore_index=True,
        ).dropna().drop_duplicates().tolist()
        has_missing_category = (
            returned_rows[category_column].isna().any()
            or delivered_rows[category_column].isna().any()
        )
        category_keys = observed_categories + (
            [None] if has_missing_category else []
        )
        for category in category_keys:
            if category is None:
                returned_count = int(returned_rows[category_column].isna().sum())
                delivered_count = int(delivered_rows[category_column].isna().sum())
            else:
                returned_count = int(returned_rows[category_column].eq(category).sum())
                delivered_count = int(delivered_rows[category_column].eq(category).sum())
            returned_share = returned_count / len(returned_rows)
            delivered_share = delivered_count / len(delivered_rows)
            category_label = "（缺失品类）" if category is None else str(category)
            category_rows.append(
                {
                    "分析分区": "品类分布（按记录统计）",
                    "指标或类别": category_label,
                    "退货组 Returned": returned_count,
                    "正常交付组 Delivered": delivered_count,
                    "差异/说明": (returned_share - delivered_share) * 100,
                    "_退货占比": returned_share,
                    "_交付占比": delivered_share,
                }
            )
        category_rows.sort(
            key=lambda row: (
                row["退货组 Returned"],
                row["_退货占比"],
            ),
            reverse=True,
        )
        for row in category_rows:
            output_rows.append(
                {
                    "分析分区": "品类分布（按记录统计）",
                    "指标或类别": row["指标或类别"],
                    "退货组 Returned": row["退货组 Returned"],
                    "正常交付组 Delivered": row["正常交付组 Delivered"],
                    "差异/说明": "记录数；同品类组内占比及百分点差异见下一行。",
                }
            )
            output_rows.append(
                {
                    "分析分区": "品类分布（按记录统计）",
                    "指标或类别": f"{row['指标或类别']}（组内占比）",
                    "退货组 Returned": row["_退货占比"],
                    "正常交付组 Delivered": row["_交付占比"],
                    "差异/说明": row["差异/说明"],
                }
            )
        add_row(
            "品类分布（按记录统计）",
            "各组品类占比口径",
            None,
            None,
            "占比分母分别为退货组/交付组全部记录（含缺失品类）；"
            "差异=退货组占比-交付组占比，单位为百分点。",
        )
        duplicate_order_rows = (
            int(returned_rows[order_column].duplicated().sum())
            if order_column is not None else None
        )
        duplicate_delivered_orders = (
            int(delivered_rows[order_column].duplicated().sum())
            if order_column is not None else None
        )
        if (duplicate_order_rows or duplicate_delivered_orders):
            warnings.append(
                "同一订单编号存在多条商品记录；品类分布按记录（商品行）统计，"
                "不能解释为订单数或退货率。"
            )
        else:
            warnings.append(
                "品类分布以记录数为基础；本数据中两组的订单编号均未重复，"
                "但结果仍表述为记录占比，不将品类占比解释为退货率。"
            )
    else:
        reason = (
            "未能唯一识别明确的品类字段。"
            if category_column is None
            else "缺少任一比较组样本，无法比较品类分布。"
        )
        add_row("品类分布（按记录统计）", "分析状态", None, None, reason)
        warnings.append(reason)

    add_section("简要观察（仅描述，不作因果推断）")
    observation_count = 0
    if compared_records:
        add_row(
            "简要观察",
            "样本量",
            len(returned_rows),
            len(delivered_rows),
            "状态组内记录数；已排除其他状态。",
        )
        observation_count += 1

    price_rows = [
        row for row in output_rows
        if row.get("分析分区") == "单价分布"
        and row.get("指标或类别") == "中位数"
    ]
    if price_rows and price_rows[0]["退货组 Returned"] is not None and price_rows[0]["正常交付组 Delivered"] is not None:
        returned_median = price_rows[0]["退货组 Returned"]
        delivered_median = price_rows[0]["正常交付组 Delivered"]
        relation = "高" if returned_median > delivered_median else (
            "低" if returned_median < delivered_median else "相同"
        )
        add_row(
            "简要观察",
            "单价中位数对比",
            returned_median,
            delivered_median,
            f"退货组中位数{relation}；差值 {returned_median - delivered_median:.2f}，"
            "仅为描述性观察。",
        )
        observation_count += 1

    discount_rows = [
        row for row in output_rows
        if row.get("分析分区") == "折扣分布"
        and row.get("指标或类别") == "中位数"
    ]
    if discount_rows and discount_rows[0]["退货组 Returned"] is not None and discount_rows[0]["正常交付组 Delivered"] is not None:
        returned_median = discount_rows[0]["退货组 Returned"]
        delivered_median = discount_rows[0]["正常交付组 Delivered"]
        relation = "高" if returned_median > delivered_median else (
            "低" if returned_median < delivered_median else "相同"
        )
        add_row(
            "简要观察",
            "折扣中位数对比（原始字段值）",
            returned_median,
            delivered_median,
            f"退货组中位数{relation}；不将折扣字段解释为金额或直接换算百分比。",
        )
        observation_count += 1

    if category_rows:
        top_category = category_rows[0]
        add_row(
            "简要观察",
            "退货组记录数最多的品类",
            top_category["指标或类别"],
            None,
            f"退货组 {top_category['_退货占比']:.1%}，"
            f"正常交付组 {top_category['_交付占比']:.1%}，"
            f"占比差异 {top_category['差异/说明']:.2f} 个百分点；"
            "该差异不是品类退货率。",
        )
        observation_count += 1

    if observation_count < 5:
        add_row(
            "简要观察",
            "解释限制",
            None,
            None,
            "仅作描述性比较，不说明变量导致退货；单价/折扣定义、订单粒度"
            "或缺失字段的限制详见以上提示。",
        )

    return pd.DataFrame(output_rows, columns=output_columns), list(dict.fromkeys(warnings))


def update_excel_sheet(
    workbook_path: Path,
    sheet_name: str,
    report: pd.DataFrame,
) -> None:
    """仅新增或替换指定工作表，保留工作簿其余工作表内容。"""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "更新 Excel 报告需要 openpyxl；请运行：python3 -m pip install -r requirements.txt"
        ) from error

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = Workbook()
        if workbook.active is not None:
            workbook.remove(workbook.active)

    old_index = None
    if sheet_name in workbook.sheetnames:
        old_index = workbook.sheetnames.index(sheet_name)
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(
        title=sheet_name,
        index=old_index if old_index is not None else len(workbook.sheetnames),
    )

    worksheet.append(list(report.columns))
    for values in report.itertuples(index=False, name=None):
        worksheet.append(
            [
                (
                    None
                    if isinstance(value, (float, int))
                    and not math.isfinite(float(value))
                    else value.item() if hasattr(value, "item") else value
                )
                for value in values
            ]
        )

    worksheet.freeze_panes = "A2"
    if worksheet.max_row > 0 and worksheet.max_column > 0:
        worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")

    section_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for row in range(2, worksheet.max_row + 1):
        section_value = worksheet.cell(row, 1).value
        is_section_header = (
            section_value is not None
            and worksheet.cell(row, 2).value is None
            and all(
                worksheet.cell(row, column).value is None
                for column in range(3, worksheet.max_column + 1)
            )
        )
        if is_section_header:
            for cell in worksheet[row]:
                cell.font = Font(bold=True, color="1F1F1F")
                cell.fill = section_fill
        for column in range(3, worksheet.max_column + 1):
            cell = worksheet.cell(row, column)
            if not isinstance(cell.value, (int, float)):
                continue
            metric = str(worksheet.cell(row, 2).value or "")
            section = str(worksheet.cell(row, 1).value or "")
            if "占比" in metric and "百分点" not in metric:
                cell.number_format = (
                    "0.00%"
                    if column in (3, 4)
                    else "0.00"
                )
            elif "百分点" in metric or (
                section.startswith("品类分布") and column == 5
            ):
                cell.number_format = "0.00"
            elif any(
                token in metric
                for token in (
                    "均值", "中位数", "标准差", "最小值",
                    "分位数", "最大值",
                )
            ):
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
            else:
                cell.number_format = "#,##0;[Red]-#,##0"

    for column_cells in worksheet.iter_cols():
        letter = column_cells[0].column_letter
        width = max(
            (
                len(str(cell.value))
                for cell in column_cells[: min(worksheet.max_row, 100)]
                if cell.value is not None
            ),
            default=10,
        )
        worksheet.column_dimensions[letter].width = min(max(width + 2, 12), 52)

    workbook.save(workbook_path)


def build_grouped_return_rates(
    data: pd.DataFrame,
) -> tuple[pd.DataFrame, list[str]]:
    """仅按 Returned/Delivered 的去重订单数计算分组退货率。"""
    warnings: list[str] = []
    output_columns = [
        "分组维度",
        "分组区间/类别",
        "有效订单总数",
        "退货订单数",
        "退货率",
        "业务注释",
    ]
    rows: list[dict] = []
    status_column, status_candidates = find_single_column(
        data, {"orderstatus", "status"}
    )
    order_column, order_candidates = find_single_column(
        data, {"orderid", "ordernumber", "transactionid"}
    )

    business_note = (
        "退货率=该组 Returned 去重订单数/该组 Returned 与 Delivered "
        "去重订单并集；Cancelled、Shipped、Pending及其他状态不纳入分母。"
        "组内占比 ≠ 退货率：组内占比是某类在退货组或交付组中的构成比例，"
        "退货率则以该分组内 Returned+Delivered 的有效订单为分母。"
        "若一个订单跨多个品类、品牌、价格档或折扣档，会分别计入对应组，"
        "不可跨组相加；所有统计基于原始记录，不修改原始数据。"
    )

    def add_unavailable(dimension: str, reason: str) -> None:
        rows.append(
            {
                "分组维度": dimension,
                "分组区间/类别": "未生成",
                "有效订单总数": None,
                "退货订单数": None,
                "退货率": None,
                "业务注释": reason,
            }
        )
        warnings.append(f"{dimension}：{reason}")

    if status_column is None:
        detail = (
            f"状态字段候选不唯一（{', '.join(map(str, status_candidates))}）。"
            if status_candidates
            else "未找到明确的订单状态字段。"
        )
        for dimension in ("Category", "Brand", "价格区间", "折扣区间"):
            add_unavailable(dimension, detail + "无法筛选 Returned/Delivered，未计算退货率。")
        return pd.DataFrame(rows, columns=output_columns), list(dict.fromkeys(warnings))

    if order_column is None:
        detail = (
            f"订单编号字段候选不唯一（{', '.join(map(str, order_candidates))}）。"
            if order_candidates
            else "未找到明确的订单编号字段。"
        )
        for dimension in ("Category", "Brand", "价格区间", "折扣区间"):
            add_unavailable(dimension, detail + "无法按订单去重，未计算退货率。")
        return pd.DataFrame(rows, columns=output_columns), list(dict.fromkeys(warnings))

    status_values = data[status_column].map(
        lambda value: str(value).strip().casefold() if pd.notna(value) else ""
    )
    returned_mask = status_values.eq("returned")
    delivered_mask = status_values.eq("delivered")
    included_mask = returned_mask | delivered_mask
    included = data.loc[included_mask].copy()
    included["_退货分组状态"] = status_values.loc[included_mask]
    # 订单号只在临时副本中去除首尾空格；原始 CSV 和 DataFrame 不会被修改。
    included["_分析订单编号"] = (
        included[order_column].astype("string").str.strip().replace("", pd.NA)
    )
    missing_order_count = int(included["_分析订单编号"].isna().sum())
    included = included.loc[included["_分析订单编号"].notna()].copy()
    if missing_order_count:
        warnings.append(
            f"Returned/Delivered 记录中有 {missing_order_count} 条订单编号缺失或为空，"
            "无法计入去重订单数，已排除于计算。"
        )

    # 同一订单若在两个状态中都出现，分母按订单编号并集去重、分子按 Returned 去重。
    status_counts_by_order = included.groupby("_分析订单编号")[
        "_退货分组状态"
    ].nunique()
    conflicting_orders = int((status_counts_by_order > 1).sum())
    if conflicting_orders:
        warnings.append(
            f"有 {conflicting_orders} 个订单编号同时出现在 Returned 和 Delivered；"
            "分母按编号并集去重，退货分子按 Returned 编号去重。"
        )

    excluded_statuses = sorted(
        {
            str(value).strip()
            for value in data.loc[~included_mask, status_column].dropna().unique()
        }
    )
    excluded_rows = int((~included_mask).sum())
    warnings.append(
        f"状态筛选仅保留 Returned 与 Delivered；排除了 {excluded_rows} 条其他状态记录"
        + (f"（{', '.join(excluded_statuses)}）。" if excluded_statuses else "。")
    )

    def append_dimension(
        dimension: str,
        group_values: pd.Series,
        expected_groups: list[str] | None = None,
        note: str = "",
    ) -> None:
        group_values = group_values.reset_index(drop=True)
        group_records = included.reset_index(drop=True).copy()
        group_records["_分组值"] = group_values
        if expected_groups is None:
            group_names = [
                str(value)
                for value in group_records["_分组值"].dropna().drop_duplicates().tolist()
            ]
            if group_records["_分组值"].isna().any():
                group_names.append("（缺失值）")
        else:
            group_names = expected_groups

        for group_name in group_names:
            if group_name == "（缺失值）":
                group = group_records.loc[group_records["_分组值"].isna()]
            else:
                group = group_records.loc[
                    group_records["_分组值"].astype("string").eq(group_name)
                ]
            order_ids = group["_分析订单编号"]
            valid_orders = int(order_ids.nunique())
            returned_orders = int(
                group.loc[
                    group["_退货分组状态"].eq("returned"),
                    "_分析订单编号",
                ].nunique()
            )
            rate = returned_orders / valid_orders if valid_orders else None
            rows.append(
                {
                    "分组维度": dimension,
                    "分组区间/类别": group_name,
                    "有效订单总数": valid_orders,
                    "退货订单数": returned_orders,
                    "退货率": rate,
                    "业务注释": note or business_note,
                }
            )

    # 明确字段别名必须唯一匹配，避免从含义不清或重名字段推算业务指标。
    category_column, category_candidates = find_single_column(
        data, {"category", "productcategory"}
    )
    brand_column, brand_candidates = find_single_column(data, {"brand"})
    if category_column is None:
        reason = (
            f"品类字段候选不唯一：{', '.join(map(str, category_candidates))}。"
            if category_candidates
            else "未找到明确的品类字段。"
        )
        add_unavailable("Category", reason)
    else:
        category_values = included[category_column].astype("string").fillna("（缺失值）")
        append_dimension(
            "Category",
            category_values,
            note=f"按字段“{category_column}”分组；" + business_note,
        )

    if brand_column is None:
        reason = (
            f"品牌字段候选不唯一：{', '.join(map(str, brand_candidates))}。"
            if brand_candidates
            else "未找到明确的品牌字段。"
        )
        add_unavailable("Brand", reason)
    else:
        brand_values = included[brand_column].astype("string").fillna("（缺失值）")
        append_dimension(
            "Brand",
            brand_values,
            note=f"按字段“{brand_column}”分组；" + business_note,
        )

    price_column, price_candidates = find_single_column(
        data, {"unitprice", "sellingprice", "price"}
    )
    if price_column is None:
        reason = (
            f"单价字段候选不唯一：{', '.join(map(str, price_candidates))}。"
            if price_candidates
            else "未找到明确的单价字段。"
        )
        add_unavailable("价格区间", reason)
    elif not pd.api.types.is_numeric_dtype(data[price_column]):
        add_unavailable(
            "价格区间",
            f"字段“{price_column}”不是数值类型，未擅自转换后分箱。",
        )
    else:
        price_values = pd.to_numeric(included[price_column], errors="coerce")
        finite_mask = price_values.map(
            lambda value: pd.isna(value) or math.isfinite(float(value))
        )
        out_of_range = price_values.notna() & (
            ~finite_mask | price_values.lt(0)
        )
        invalid_count = int((price_values.isna() | out_of_range).sum())
        if invalid_count:
            warnings.append(
                f"UnitPrice 分组中有 {invalid_count} 条单价缺失、非有限或小于 0，"
                "不属于给定区间，未计入价格分箱。"
            )
        price_labels = [
            "[0, 100)",
            "[100, 200)",
            "[200, 300)",
            "[300, 500)",
            "[500, ∞)",
        ]
        price_bins = [0, 100, 200, 300, 500, float("inf")]
        valid_price = price_values.where(finite_mask & price_values.ge(0))
        price_groups = pd.cut(
            valid_price,
            bins=price_bins,
            labels=price_labels,
            right=False,
            include_lowest=True,
        )
        append_dimension(
            "价格区间",
            price_groups,
            expected_groups=price_labels,
            note=f"按字段“{price_column}”分箱；区间左闭右开；"
            + business_note,
        )

    discount_column, discount_candidates = find_single_column(
        data, {"discount", "discountrate", "discountpercentage", "discountpercent"}
    )
    if discount_column is None:
        reason = (
            f"折扣字段候选不唯一：{', '.join(map(str, discount_candidates))}。"
            if discount_candidates
            else "未找到明确的折扣字段。"
        )
        add_unavailable("折扣区间", reason)
    elif not pd.api.types.is_numeric_dtype(data[discount_column]):
        add_unavailable(
            "折扣区间",
            f"字段“{discount_column}”不是数值类型，未擅自转换后分箱。",
        )
    else:
        discount_values = pd.to_numeric(included[discount_column], errors="coerce")
        finite_mask = discount_values.map(
            lambda value: pd.isna(value) or math.isfinite(float(value))
        )
        in_range = (
            finite_mask
            & discount_values.notna()
            & discount_values.ge(0)
            & discount_values.le(0.3)
        )
        invalid_count = int((~in_range).sum())
        if invalid_count:
            warnings.append(
                f"Discount 分组中有 {invalid_count} 条折扣缺失、非有限或不在 "
                "[0, 0.3] 范围内，不属于给定分箱，未计入折扣区间。"
            )
        discount_labels = [
            "[0, 0.1)",
            "[0.1, 0.2)",
            "[0.2, 0.3]",
        ]
        discount_bins = [0, 0.1, 0.2, 0.3]
        discount_groups = pd.cut(
            discount_values.where(in_range),
            bins=discount_bins,
            labels=discount_labels,
            right=False,
            include_lowest=True,
        )
        # pd.cut 的右开区间默认不含最后边界；按定义将精确 0.3 纳入末档。
        discount_groups = discount_groups.astype(object)
        discount_groups.loc[discount_values.eq(0.3) & in_range] = discount_labels[-1]
        append_dimension(
            "折扣区间",
            discount_groups,
            expected_groups=discount_labels,
            note=f"按字段“{discount_column}”分箱；折扣沿用字段原值，"
            "不解释为金额；前两档左闭右开，最后一档包含 0.3。"
            + business_note,
        )

    return pd.DataFrame(rows, columns=output_columns), list(dict.fromkeys(warnings))


def update_grouped_return_rate_report(
    workbook_path: Path,
    report: pd.DataFrame,
) -> None:
    """只更新分组退货率工作表，并在业务提示中补充统计口径说明。"""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "更新 Excel 报告需要 openpyxl；请运行：python3 -m pip install -r requirements.txt"
        ) from error

    if not workbook_path.is_file():
        raise FileNotFoundError(
            f"找不到现有销售分析报告：{workbook_path}；为避免覆盖或重建其他工作表，已停止更新。"
        )

    workbook = load_workbook(workbook_path)
    sheet_name = "分组退货率"
    old_index = (
        workbook.sheetnames.index(sheet_name)
        if sheet_name in workbook.sheetnames
        else len(workbook.sheetnames)
    )
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name, index=old_index)
    worksheet.append(list(report.columns))
    for values in report.itertuples(index=False, name=None):
        worksheet.append(
            [
                (
                    None
                    if isinstance(value, (float, int))
                    and not math.isfinite(float(value))
                    else value.item() if hasattr(value, "item") else value
                )
                for value in values
            ]
        )

    worksheet.freeze_panes = "A2"
    if worksheet.max_row and worksheet.max_column:
        worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row in range(2, worksheet.max_row + 1):
        rate_cell = worksheet.cell(row, 5)
        if isinstance(rate_cell.value, (float, int)):
            rate_cell.number_format = "0.00%"
        for column in (3, 4):
            count_cell = worksheet.cell(row, column)
            if isinstance(count_cell.value, (float, int)):
                count_cell.number_format = "#,##0"

    widths = {"A": 18, "B": 24, "C": 18, "D": 18, "E": 14, "F": 72}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2):
        row[5].alignment = row[5].alignment.copy(wrap_text=True, vertical="top")

    # 将解释追加到现有业务提示页，不更改原有内容、格式、冻结窗格或筛选范围。
    business_tip = (
        "分组退货率口径：有效订单总数按分组内 Returned 与 Delivered "
        "的不同 OrderID 并集去重；退货订单数按 Returned 的不同 OrderID 去重；"
        "退货率=退货订单数/有效订单总数。Cancelled、Shipped、Pending "
        "及其他状态不参与分母。组内占比 ≠ 退货率：组内占比描述某组别在状态组中的构成，"
        "退货率描述某分组订单中退货订单的比例。一个订单跨多个品类、品牌或区间时会出现在多个组，"
        "组间结果不可相加。分析只读取原始数据，不修改原始数据。"
    )
    if "业务提示" in workbook.sheetnames:
        tips = workbook["业务提示"]
        existing_tips = {
            str(tips.cell(row, 1).value)
            for row in range(2, tips.max_row + 1)
            if tips.cell(row, 1).value is not None
        }
        if business_tip not in existing_tips:
            next_row = tips.max_row + 1
            tips.cell(next_row, 1, business_tip)
            if next_row > 2:
                previous = tips.cell(next_row - 1, 1)
                current = tips.cell(next_row, 1)
                if previous.has_style:
                    current._style = previous._style
                if previous.number_format:
                    current.number_format = previous.number_format
                if previous.alignment:
                    current.alignment = previous.alignment.copy()
                if previous.protection:
                    current.protection = previous.protection.copy()
            if tips.row_dimensions[next_row - 1].height:
                tips.row_dimensions[next_row].height = tips.row_dimensions[next_row - 1].height
    else:
        tips = workbook.create_sheet("业务提示")
        tips.append(["提示"])
        tips.append([business_tip])
        tips.freeze_panes = "A2"
        tips.auto_filter.ref = tips.dimensions

    workbook.save(workbook_path)


def write_formatted_excel(
    sheets: dict[str, pd.DataFrame],
    output_path: Path,
) -> None:
    """写入带冻结首行、筛选和易读数字格式的 Excel 工作簿。"""
    try:
        from openpyxl.styles import Font, PatternFill
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "生成 Excel 报告需要 openpyxl；请运行：python3 -m pip install -r requirements.txt"
        ) from error

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, report in sheets.items():
            report.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            if worksheet.max_row > 0 and worksheet.max_column > 0:
                worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(
                    name=cell.font.name,
                    size=cell.font.sz,
                    bold=True,
                    color="FFFFFF",
                )
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="1F4E78",
                )

            # 按列标题格式化数值，同时避免将计数和销量显示为小数。
            for column_cells in worksheet.iter_cols(min_row=2):
                header = str(worksheet.cell(1, column_cells[0].column).value or "")
                for cell in column_cells:
                    if not isinstance(cell.value, (int, float)):
                        continue
                    is_amount_or_stat = any(
                        token in header
                        for token in (
                            "金额", "单价", "价格", "均值", "中位数", "标准差",
                            "最小值", "最大值", "百分位", "下界", "上界",
                        )
                    )
                    if isinstance(cell.value, float) or is_amount_or_stat:
                        cell.number_format = "#,##0.00;[Red]-#,##0.00"
                    else:
                        cell.number_format = "#,##0;[Red]-#,##0"

            # 基于表头及有限样例调整列宽，限制过长文本造成的横向膨胀。
            for column_cells in worksheet.iter_cols():
                column_letter = column_cells[0].column_letter
                width = max(
                    (len(str(cell.value)) for cell in column_cells[:51] if cell.value is not None),
                    default=10,
                )
                worksheet.column_dimensions[column_letter].width = min(max(width + 2, 10), 48)


if __name__ == "__main__":
    main()
